#!/usr/bin/env python3
"""UltimaWeapon — neutral workbook builder (openpyxl). One JSON spec in, one brand-free .xlsx out.

Spec (JSON):
{
  "title": "Full SEO, Compliance & Competitive Audit",          # goes in the READ ME tab and the file's core Title only
  "readme": ["Prepared for Example Co.", "Tier 1 (GSC + GA4 exports)", "..."],   # optional lines -> READ ME tab (first)
  "tabs": [
    {"name": "FINDINGS", "rows": [{"Module": "Technical", "Severity": "High", "Finding": "...", "URL": "https://..."}, ...],
     "columns": ["Module", "Severity", "Finding", "URL"],          # optional order; default = first-row keys + union of the rest
     "widths": {"Finding": 70}},                                    # optional per-column width override
    {"name": "Verify Queue", "rows": [...]}
  ]
}
Rules baked in: frozen header row, autofilter, wrapped text, live hyperlinks for any cell whose value starts with http,
green / amber / red fills for Severity, Priority, Disposition, Status, Tier and Reading columns, a neutral palette,
no placeholder rows (empty tabs get one line saying "No rows — see Data Sources & Coverage"), and document properties
set to a neutral title with blank creator / lastModifiedBy. Sheet names are trimmed to Excel's 31 chars.

Usage
  python3 build_workbook.py spec.json out.xlsx [--creator "Recipient Name"]
"""
import argparse, datetime, json, re, sys
try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("openpyxl missing: pip install openpyxl")

HEADER_FILL = PatternFill("solid", fgColor="2F3B4A"); HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
BODY_FONT = Font(name="Calibri", size=10); LINK_FONT = Font(name="Calibri", size=10, color="1F5FBF", underline="single")
THIN = Side(style="thin", color="D9DEE3"); BORDER = Border(top=THIN, bottom=THIN, left=THIN, right=THIN)
GREEN, AMBER, RED, GREY = (PatternFill("solid", fgColor=c) for c in ("E3F1E4", "FFF1D6", "F9DCDC", "EEF0F2"))
STATUS_COLS = re.compile(r"^(severity|priority|disposition|status|tier|reading|verdict|confidence|grade|bucket|systemic\??|match)$", re.I)
def status_fill(v):
    s = str(v).strip().lower()
    if re.search(r"^(high|critical|s1|s2|confirmed|red|fail|poor|disavow|match|f|d)\b|blocker|attack|urgent|not_tested|not tested|waf", s): return RED
    if re.search(r"^(medium|s3|candidate|amber|needs|review|plausible|c|watch|fingerprint|none)\b|partial|unverified|verify", s): return AMBER
    if re.search(r"^(low|s4|cleared|green|pass|good|keep|ok|a|b|confirmed-ok|raw|rendered|consistent|win)\b|working|strength", s): return GREEN
    return None

def add_tab(wb, name, rows, columns=None, widths=None):
    ws = wb.create_sheet(title=re.sub(r"[\[\]\*\?/\\:]", "-", name)[:31])
    if not rows:
        ws["A1"] = "No rows — see Data Sources & Coverage"; ws["A1"].font = BODY_FONT; return ws
    cols = list(columns or [])
    for r in rows:
        for k in r.keys():
            if k not in cols: cols.append(k)
    ws.append(cols)
    for c in range(1, len(cols) + 1):
        cell = ws.cell(row=1, column=c); cell.fill = HEADER_FILL; cell.font = HEADER_FONT; cell.alignment = Alignment(vertical="center", wrap_text=True); cell.border = BORDER
    for r in rows:
        ws.append([("" if r.get(c) is None else (json.dumps(r[c], ensure_ascii=False) if isinstance(r[c], (dict, list)) else r[c])) for c in cols])
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(cols)):
        for cell in row:
            cell.font = BODY_FONT; cell.border = BORDER; cell.alignment = Alignment(vertical="top", wrap_text=True)
            v = cell.value
            if isinstance(v, str) and re.match(r"https?://\S+$", v.strip()):
                cell.hyperlink = v.strip(); cell.font = LINK_FONT
            hdr = cols[cell.column - 1]
            if STATUS_COLS.match(str(hdr)) and v not in (None, ""):
                f = status_fill(v)
                if f: cell.fill = f
    for i, c in enumerate(cols, start=1):
        longest = max([len(str(c))] + [len(str(r.get(c, ""))) for r in rows[:300]])
        w = (widths or {}).get(c) or min(max(12, int(longest * 0.9) + 2), 60)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    return ws

def build(spec, out, creator=""):
    wb = Workbook(); wb.remove(wb.active)
    readme = spec.get("readme") or []
    ws = wb.create_sheet("READ ME")
    ws["A1"] = spec.get("title", "Audit Workbook"); ws["A1"].font = Font(bold=True, size=14, name="Calibri")
    ws["A2"] = f"Generated {datetime.date.today().isoformat()}"; ws["A2"].font = BODY_FONT
    for i, line in enumerate(readme, start=4):
        ws.cell(row=i, column=1, value=line).font = BODY_FONT; ws.cell(row=i, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 120
    for t in spec.get("tabs", []):
        add_tab(wb, t["name"], t.get("rows", []), t.get("columns"), t.get("widths"))
    p = wb.properties; p.title = spec.get("title", "Audit Workbook"); p.creator = creator or ""; p.lastModifiedBy = creator or ""
    p.subject = ""; p.description = ""; p.keywords = ""; p.category = ""; p.company = None if hasattr(p, "company") else None
    wb.save(out); print(f"wrote {out}: {len(spec.get('tabs', []))} tabs + READ ME")

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("spec"); ap.add_argument("out"); ap.add_argument("--creator", default="")
    a = ap.parse_args(); build(json.load(open(a.spec)), a.out, a.creator)

if __name__ == "__main__": main()
