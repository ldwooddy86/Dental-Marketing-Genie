#!/usr/bin/env python3
"""
legaldocx.py — mechanical clean of a .docx in place, plus text extraction.

Editing run text in place is what preserves the client's formatting: bold,
headings, tables, hyperlinks, and styles all survive because we never rebuild
the document, we only rewrite the strings inside existing runs.

  python3 legaldocx.py clean INPUT.docx -o OUTPUT.docx --text-out extracted.txt
  python3 legaldocx.py extract INPUT.docx --text-out extracted.txt

After cleaning, run legalscan.py analyze/scan on the extracted text to score
readability and build the legal verification worklist. Apply judgment-level
edits back into the .docx with the docx skill's tooling.

A note on flags: this script deliberately does NOT insert flag text into the
document. Flags belong in the change log, where an attorney reads them, not
scattered through copy a marketer may paste straight to a CMS.
"""

import argparse
import os
import re
import sys

try:
    import docx  # python-docx
except ImportError:
    print("python-docx is required:  pip install python-docx --break-system-packages",
          file=sys.stderr)
    raise

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from legalscan import INVISIBLES, SMART_MAP, FILLER_LEADS  # noqa: E402


def clean_run_text(t, emdash="comma", quotes="smart", counts=None):
    """Same deterministic rules as legalscan.clean_text, run-safe.

    Word documents are usually a print-or-CMS deliverable, so smart quotes are
    kept by default here (unlike the plain-text path, which normalizes them).
    """
    c = counts if counts is not None else {}

    def bump(k, n=1):
        c[k] = c.get(k, 0) + n

    for bad, good in INVISIBLES.items():
        if bad in t:
            bump("invisible_unicode", t.count(bad))
            t = t.replace(bad, good)

    if quotes == "straight":
        for bad, good in SMART_MAP.items():
            if bad in t:
                bump("smart_quotes", t.count(bad))
                t = t.replace(bad, good)

    if emdash != "keep":
        repl = ", " if emdash == "comma" else " - "
        n = len(re.findall(r"\s*—\s*", t))
        if n:
            bump("em_dashes", n)
            t = re.sub(r"\s*—\s*", repl, t)
        t = re.sub(r"(?<=\d)\s*–\s*(?=\d)", "-", t)
        n = len(re.findall(r"\s*–\s*", t))
        if n:
            bump("en_dashes", n)
            t = re.sub(r"\s*–\s*", repl, t)
        n = len(re.findall(r"\s*--\s*", t))
        if n:
            bump("double_hyphen", n)
            t = re.sub(r"\s*--\s*", repl, t)

    if "…" in t:
        bump("ellipsis_char", t.count("…"))
        t = t.replace("…", "...")

    def amp(m):
        bump("ampersand")
        return f"{m.group(1)} and {m.group(2)}"
    t = re.sub(r"\b([a-z]+) & ([a-z]+)\b", amp, t)

    for pat, r in FILLER_LEADS:
        new, n = re.subn(pat, r, t, flags=re.I)
        if n:
            bump("filler", n)
            t = new

    new, n = re.subn(r"\b(the|a|an|to|of|and|is|in|for|that|with) \1\b", r"\1",
                     t, flags=re.I)
    if n:
        bump("duplicate_words", n)
        t = new

    new, n = re.subn(r" {2,}", " ", t)
    if n:
        bump("double_spaces", n)
        t = new

    return t, c


def iter_paragraphs(document):
    """Body paragraphs plus everything inside tables, headers, and footers."""
    for p in document.paragraphs:
        yield p
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p
    for section in document.sections:
        for part in (section.header, section.footer,
                     section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            if part is None:
                continue
            for p in part.paragraphs:
                yield p


def extract_text(document):
    """Flatten to text, keeping list markers so bullet density stays measurable."""
    lines = []
    for p in document.paragraphs:
        style = (p.style.name or "").lower()
        txt = p.text
        if not txt.strip():
            lines.append("")
            continue
        if "list" in style or "bullet" in style:
            lines.append(f"- {txt}")
        elif style.startswith("heading"):
            level = re.search(r"(\d)", style)
            hashes = "#" * (int(level.group(1)) if level else 2)
            lines.append(f"\n{hashes} {txt}")
        else:
            lines.append(txt)
    for table in document.tables:
        lines.append("")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("clean", "extract"):
        p = sub.add_parser(name)
        p.add_argument("input")
        p.add_argument("-o", "--output")
        p.add_argument("--text-out")
        p.add_argument("--emdash", choices=["comma", "hyphen", "keep"], default="comma")
        p.add_argument("--quotes", choices=["smart", "straight", "keep"], default="smart")
    args = ap.parse_args()

    document = docx.Document(args.input)

    if args.cmd == "clean":
        counts, touched = {}, 0
        for p in iter_paragraphs(document):
            for run in p.runs:
                if not run.text:
                    continue
                new, counts = clean_run_text(run.text, args.emdash,
                                             args.quotes, counts)
                if new != run.text:
                    run.text = new
                    touched += 1
        out = args.output or (os.path.splitext(args.input)[0] + "_cleaned.docx")
        document.save(out)
        print(f"Cleaned {touched} runs -> {out}")
        if counts:
            for k, v in sorted(counts.items(), key=lambda kv: -kv[1]):
                print(f"  {v:>4}  {k}")
        else:
            print("  (no mechanical artifacts found)")

    if args.text_out:
        with open(args.text_out, "w", encoding="utf-8") as fh:
            fh.write(extract_text(document))
        print(f"Text extraction -> {args.text_out}")
        print("  Next: python3 legalscan.py all "
              f"{args.text_out} --json report.json")

    return 0


if __name__ == "__main__":
    sys.exit(main())
